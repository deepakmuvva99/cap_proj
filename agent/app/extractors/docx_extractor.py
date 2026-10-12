"""DOCX Extractor for Microsoft Word invoice documents."""
import logging
from pathlib import Path
from typing import Dict, Any, Optional
import docx
from app.extractors.parser_utils import extract_and_merge_fields

logger = logging.getLogger("agent.extractors.docx")


def extract_from_docx(file_path: str, meta_data: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
    """Extract raw text and parsed fields from a .docx document."""
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"DOCX file not found: {file_path}")

    doc = docx.Document(path)
    paragraphs_text = [p.text for p in doc.paragraphs if p.text.strip()]

    # Extract tables content
    table_lines = []
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                table_lines.append(" | ".join(row_cells))

    full_text = "\n".join(paragraphs_text + table_lines).strip()
    parsed_fields = extract_and_merge_fields(full_text, meta_data)

    return {
        "raw_text": full_text,
        "extraction_method": "docx",
        "fields": parsed_fields,
    }
